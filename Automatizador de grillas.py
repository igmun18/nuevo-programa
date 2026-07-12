import sys
import os
import urllib.request
import json
import threading
import customtkinter as ctk
from tkinterdnd2 import DND_FILES, TkinterDnD
from tkinter import filedialog, messagebox
from PIL import Image, ImageTk
import openpyxl
import re
from openpyxl import load_workbook
from copy import copy
from dateutil.relativedelta import relativedelta
from datetime import datetime
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment


def resource_path(filename):
    """Devuelve la ruta correcta tanto en desarrollo como en .exe"""
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, filename)
    return os.path.join(os.path.abspath(os.path.dirname(__file__)), filename)


# ================================================================
# KILL SWITCH — URL del Gist raw
# ================================================================
GIST_URL = "https://gist.githubusercontent.com/igmun18/b6f9d93790085925de85db541f67e32e/raw/licencia.json"

def verificar_licencia():
    try:
        import random
        url = f"{GIST_URL}?nocache={random.randint(1, 999999)}"
        req = urllib.request.urlopen(url, timeout=5)
        data = json.loads(req.read().decode())
        return data.get("activo", False)
    except urllib.error.URLError:
        return False
    except Exception:
        return False


# ================================================================
# SOPORTE XLS LEGACY
# ================================================================

def cargar_workbook_cualquier_formato(ruta):
    extension = os.path.splitext(ruta)[1].lower()

    if extension == ".xlsx":
        return openpyxl.load_workbook(ruta, data_only=True)

    elif extension == ".xls":
        try:
            import xlrd
        except ImportError:
            raise ImportError(
                "Para abrir archivos .xls instalá xlrd:\n"
                "  pip install xlrd\n"
                "O convertí el archivo a .xlsx desde Excel."
            )

        xls_wb = xlrd.open_workbook(ruta, formatting_info=True)
        new_wb = openpyxl.Workbook()
        new_wb.remove(new_wb.active)

        for sheet_name in xls_wb.sheet_names():
            xls_ws = xls_wb.sheet_by_name(sheet_name)
            new_ws = new_wb.create_sheet(title=sheet_name)

            for row_idx in range(xls_ws.nrows):
                for col_idx in range(xls_ws.ncols):
                    cell = xls_ws.cell(row_idx, col_idx)
                    ctype = cell.ctype

                    if ctype == xlrd.XL_CELL_DATE:
                        try:
                            dt_tuple = xlrd.xldate_as_tuple(cell.value, xls_wb.datemode)
                            valor = None if dt_tuple[0] == 0 else datetime(*dt_tuple)
                        except Exception:
                            valor = None
                    elif ctype == xlrd.XL_CELL_ERROR:
                        valor = None
                    else:
                        valor = cell.value

                    new_ws.cell(row=row_idx + 1, column=col_idx + 1).value = valor

        return new_wb

    else:
        raise ValueError(f"Formato no soportado: {extension}")


def obtener_rango_combinado(ws, fila, columna):
    if fila is None or columna is None:
        return None

    for rango in ws.merged_cells.ranges:
        try:
            min_row, max_row = rango.min_row, rango.max_row
            min_col, max_col = rango.min_col, rango.max_col

            if None in (min_row, max_row, min_col, max_col):
                continue

            if (min_row <= fila <= max_row) and (min_col <= columna <= max_col):
                return rango
        except Exception:
            continue

    return None


# ================================================================
# PANTALLA DE CARGA CON GIF ANIMADO
# ================================================================

class LoadingScreen:
    """
    Ventana modal que muestra un GIF animado mientras se procesa.
    Se maneja desde el hilo principal con after() para animar el GIF.
    """

    def __init__(self, parent, gif_path=None):
        self.parent = parent
        self.ventana = ctk.CTkToplevel(parent)
        self.ventana.configure(fg_color="black") 
        self.ventana.title("Procesando...")
        self.ventana.geometry("320x260")
        self.ventana.resizable(False, False)
        self.ventana.grab_set()
        self.ventana.protocol("WM_DELETE_WINDOW", lambda: None)  # bloquear cierre manual

        # centrar sobre la ventana principal
        self.ventana.update_idletasks()
        px = parent.winfo_x() + (parent.winfo_width() // 2) - 160
        py = parent.winfo_y() + (parent.winfo_height() // 2) - 130
        self.ventana.geometry(f"+{px}+{py}")

        # intentar aplicar ícono
        try:
            self.ventana.after(200, lambda: self.ventana.wm_iconbitmap(resource_path("icono.ico")))
        except Exception:
            pass

        self.frames = []
        self.frame_idx = 0
        self._after_id = None
        self._gif_label = None

        # --- GIF ---
        if gif_path and os.path.exists(gif_path):
            self._cargar_gif(gif_path)
        else:
            # fallback: spinner de texto
            self._label_fallback = ctk.CTkLabel(
                self.ventana,
                text="⏳",
                font=("Arial", 48)
            )
            self._label_fallback.pack(pady=30)

        # --- Texto ---
        self.label_texto = ctk.CTkLabel(
            self.ventana,
            text="Procesando archivos...",
            font=("Arial", 14)
        )
        self.label_texto.pack(pady=10)

        self.label_sub = ctk.CTkLabel(
            self.ventana,
            text="Por favor esperá",
            font=("Arial", 11),
            text_color="gray"
        )
        self.label_sub.pack()

        # barra de progreso indeterminada
        self.barra = ctk.CTkProgressBar(self.ventana, mode="indeterminate", width=260)
        self.barra.pack(pady=15)
        self.barra.start()

        self.ventana.update()

    def _cargar_gif(self, gif_path):
        """Extrae los frames del GIF y los guarda como PhotoImage."""
        try:
            gif = Image.open(gif_path)
            self.frames = []

            try:
                while True:
                    frame = gif.copy().convert("RGBA").resize((200, 150), Image.LANCZOS)
                    self.frames.append(ImageTk.PhotoImage(frame))
                    gif.seek(gif.tell() + 1)
            except EOFError:
                pass

            if self.frames:
                self._gif_label = ctk.CTkLabel(self.ventana, text="")
                self._gif_label.pack(pady=10)
                self._animar()

        except Exception as e:
            print(f"Error cargando GIF: {e}")

    def _animar(self):
        """Avanza un frame y se re-agenda."""
        if not self.frames or self._gif_label is None:
            return
        try:
            self._gif_label.configure(image=self.frames[self.frame_idx])
            self.frame_idx = (self.frame_idx + 1) % len(self.frames)
            self._after_id = self.ventana.after(50, self._animar)  # ~20 fps
        except Exception:
            pass

    def cerrar(self):
        """Detiene la animación y destruye la ventana."""
        if self._after_id:
            try:
                self.ventana.after_cancel(self._after_id)
            except Exception:
                pass
        try:
            self.barra.stop()
            self.ventana.grab_release()
            self.ventana.destroy()
        except Exception:
            pass


# ================================================================
# APP PRINCIPAL
# ================================================================

class App(ctk.CTk, TkinterDnD.Tk):
    def __init__(self):
        ctk.CTk.__init__(self)
        TkinterDnD.Tk.__init__(self)

        ctk.set_appearance_mode("dark")
        self.title("Automatizador de grillas:")
        self.geometry("450x550")
        self.minsize(450, 550)
        self.configure(fg_color="#242424")

        self.after(200, self._set_icon)

        self.protocol("WM_DELETE_WINDOW", self.cerrar_programa)

        # ---------------- ICONO EXCEL ----------------
        try:
            self.icono_excel = ctk.CTkImage(
                light_image=Image.open(resource_path("excel_icon.png")),
                dark_image=Image.open(resource_path("excel_icon.png")),
                size=(40, 40)
            )
        except Exception:
            self.icono_excel = None

        self.rutas_archivos = []

        # ---------------- CONTENEDOR PRINCIPAL ----------------
        self.container = ctk.CTkFrame(self, fg_color="transparent")
        self.container.pack(fill="both", expand=True, padx=20, pady=20)

        # ---------------- SECCIÓN BUSCAR ----------------
        frame_path = ctk.CTkFrame(self.container)
        frame_path.pack(fill="x", pady=10)

        btn_buscar = ctk.CTkButton(frame_path, text="Buscar", width=100, command=self.buscar_archivo)
        btn_buscar.pack(side="left", padx=10)

        self.entry_path = ctk.CTkEntry(frame_path, placeholder_text="Ruta del archivo...")
        self.entry_path.pack(side="left", padx=10, expand=True, fill="x")

        # ---------------- SECCIÓN PORCENTAJE ----------------
        frame_porcentaje = ctk.CTkFrame(self.container)
        frame_porcentaje.pack(fill="x", pady=10)

        label = ctk.CTkLabel(frame_porcentaje, text="Porcentaje de aumento (%):")
        label.pack(side="left", padx=10)

        self.entry_porcentaje = ctk.CTkEntry(frame_porcentaje, width=100)
        self.entry_porcentaje.insert(0, "")
        self.entry_porcentaje.pack(side="left", padx=10)

        # ---------------- DRAG & DROP ----------------
        self.frame_drop = ctk.CTkFrame(
            self.container,
            border_width=2,
            border_color="#3B3B3B"
        )
        self.frame_drop.pack(fill="both", expand=True, pady=10)

        self.label_drop = ctk.CTkLabel(
            self.frame_drop,
            text="Arrastra archivos Excel aquí o usa 'Buscar'"
        )
        self.label_drop.pack(pady=10)

        self.lista_iconos_frame = ctk.CTkScrollableFrame(
            self.frame_drop,
            label_text="Archivos seleccionados"
        )
        self.lista_iconos_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # ---------------- BOTÓN ----------------
        self.btn_ejecutar = ctk.CTkButton(
            self.container,
            text="APLICAR INCREMENTO",
            fg_color="#0B1575",
            hover_color="#2808B4",
            height=45,
            command=self.ejecutar_macro
        )
        self.btn_ejecutar.pack(fill="x", pady=15)

        # DRAG & DROP
        self.frame_drop.drop_target_register(DND_FILES)
        self.frame_drop.dnd_bind('<<Drop>>', self.al_soltar_archivo)

    # ---------------- ÍCONO ----------------

    def _set_icon(self):
        try:
            ruta_ico = resource_path("icono.ico")
            self.wm_iconbitmap(ruta_ico)
            self._ruta_ico = ruta_ico
        except Exception as e:
            print(f"Error cargando icono: {e}")

    def _set_icon_toplevel(self, ventana):
        try:
            ventana.wm_iconbitmap(self._ruta_ico)
        except Exception as e:
            print(f"Error cargando icono en ventana secundaria: {e}")

    # ---------------- FUNCIONES ----------------

    def buscar_archivo(self):
        archivos = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx *.xls")])
        if archivos:
            self.agregar_a_lista(archivos)

    def al_soltar_archivo(self, event):
        rutas = self.tk.splitlist(event.data)
        self.agregar_a_lista(rutas)

    def agregar_a_lista(self, rutas):
        for ruta in rutas:
            ruta = ruta.strip('{}')
            if ruta not in self.rutas_archivos and ruta.lower().endswith(('.xlsx', '.xls')):
                self.rutas_archivos.append(ruta)
        self.renderizar_iconos()

    def confirmar_incremento(self):
        ventana = ctk.CTkToplevel(self)
        ventana.update_idletasks()
        ventana.after(200, lambda: self._set_icon_toplevel(ventana))
        ventana.title("Confirmar incremento")
        ventana.geometry("600x450")
        ventana.grab_set()

        resultado = {"aceptado": False}

        ctk.CTkLabel(
            ventana,
            text=f"¿Desea aplicar el siguiente incremento: {self.entry_porcentaje.get()}%?",
            font=("Arial", 18, "bold"),
            text_color="red"
        ).pack(pady=15)

        ctk.CTkLabel(ventana, text="Archivos seleccionados:").pack()

        frame_lista = ctk.CTkScrollableFrame(ventana, width=500, height=250)
        frame_lista.pack(fill="both", expand=True, padx=20, pady=10)

        for ruta in self.rutas_archivos:
            ctk.CTkLabel(
                frame_lista,
                text=f"✓ {os.path.basename(ruta)}",
                anchor="w"
            ).pack(fill="x", padx=10, pady=2)

        frame_botones = ctk.CTkFrame(ventana, fg_color="transparent")
        frame_botones.pack(pady=15)

        def cancelar():
            ventana.destroy()

        def aceptar():
            resultado["aceptado"] = True
            ventana.destroy()

        ctk.CTkButton(
            frame_botones,
            text="Cancelar",
            fg_color="#A12121",
            hover_color="#E63946",
            command=cancelar
        ).pack(side="left", padx=10)

        ctk.CTkButton(
            frame_botones,
            text="Aplicar incremento",
            fg_color="#0B1575",
            hover_color="#2808B4",
            command=aceptar
        ).pack(side="left", padx=10)

        self.wait_window(ventana)
        return resultado["aceptado"]

    # ---------------- GRILLA DE ICONOS ----------------

    def renderizar_iconos(self):
        for widget in self.lista_iconos_frame.winfo_children():
            widget.destroy()

        columnas = 4

        for index, ruta in enumerate(self.rutas_archivos):
            nombre = os.path.basename(ruta)
            fila = index // columnas
            col = index % columnas

            item = ctk.CTkFrame(self.lista_iconos_frame, width=120, height=120)
            item.grid(row=fila, column=col, padx=10, pady=10)
            item.grid_propagate(False)

            label_img = ctk.CTkLabel(
                item,
                image=self.icono_excel,
                text="" if self.icono_excel else "📄"
            )
            label_img.pack(pady=(10, 5))

            nombre_corto = nombre[:12] + "..." if len(nombre) > 12 else nombre
            ctk.CTkLabel(item, text=nombre_corto).pack()

            ctk.CTkButton(
                item,
                text="✕",
                width=20,
                height=20,
                fg_color="#A12121",
                hover_color="#E63946",
                command=lambda r=ruta: self.eliminar_archivo(r)
            ).place(relx=1, rely=0, anchor="ne")

    def eliminar_archivo(self, ruta):
        if ruta in self.rutas_archivos:
            self.rutas_archivos.remove(ruta)
        self.renderizar_iconos()

    def cerrar_programa(self):
        try:
            self.quit()
            self.destroy()
        except Exception:
            pass

    # ---------------- MACRO ----------------

    def ejecutar_macro(self):
        try:
            factor = 1 + (float(self.entry_porcentaje.get()) / 100)
        except ValueError:
            messagebox.showerror("Error", "Ingresa un porcentaje válido.")
            return

        if not self.rutas_archivos:
            messagebox.showwarning("Atención", "No hay archivos seleccionados.")
            return

        if not self.confirmar_incremento():
            return

        # Deshabilitar botón mientras procesa
        self.btn_ejecutar.configure(state="disabled")

        # Ruta del GIF (empaquetado junto al .exe)
        gif_path = resource_path("carga.gif")

        # Abrir pantalla de carga
        loading = LoadingScreen(self, gif_path=gif_path)

        rutas_a_procesar = list(self.rutas_archivos)

        resultado = {"exitos": 0, "error": None}

        def worker():
            """Procesamiento en hilo separado para no bloquear la UI."""
            try:
                resultado["exitos"] = self._procesar_archivos(rutas_a_procesar, factor)
            except Exception as e:
                resultado["error"] = str(e)
            finally:
                # Volver al hilo principal para cerrar la pantalla y mostrar resultado
                self.after(0, lambda: self._finalizar(loading, resultado))

        hilo = threading.Thread(target=worker, daemon=True)
        hilo.start()

    def _finalizar(self, loading, resultado):
        """Se ejecuta en el hilo principal al terminar el procesamiento."""
        loading.cerrar()
        self.btn_ejecutar.configure(state="normal")

        self.rutas_archivos.clear()
        self.entry_path.delete(0, "end")
        self.renderizar_iconos()

        if resultado["error"]:
            messagebox.showerror("Error inesperado", resultado["error"])
        else:
            messagebox.showinfo("Hecho", f"Se procesaron {resultado['exitos']} archivos correctamente.")

    def _procesar_archivos(self, rutas, factor):
        """
        Lógica de negocio pura — corre en hilo secundario.
        Retorna cantidad de archivos procesados exitosamente.
        """
        exitos = 0

        for ruta in rutas:
            try:
                wb = cargar_workbook_cualquier_formato(ruta)
                fecha_archivo = None

                for ws in wb.worksheets:
                    fila_header = None

                    # BUSCAR HEADER (fila que contiene fecha)
                    for fila in range(1, 15):
                        for col in range(ws.max_column, 0, -1):
                            valor = ws.cell(row=fila, column=col).value

                            if isinstance(valor, datetime):
                                fila_header = fila
                                break
                            elif isinstance(valor, str):
                                try:
                                    datetime.strptime(valor, "%d/%m/%Y")
                                    fila_header = fila
                                    break
                                except Exception:
                                    pass

                        if fila_header:
                            break

                    if not fila_header:
                        continue

                    # BUSCAR ÚLTIMA COLUMNA CON FECHA REAL
                    ultima_columna = None

                    for col in range(ws.max_column, 0, -1):
                        valor = ws.cell(row=fila_header, column=col).value

                        if isinstance(valor, datetime):
                            ultima_columna = col
                            break
                        elif isinstance(valor, str):
                            try:
                                datetime.strptime(valor, "%d/%m/%Y")
                                ultima_columna = col
                                break
                            except Exception:
                                pass

                    if ultima_columna is None:
                        continue

                    # DETECTAR BLOQUE REAL (merge)
                    rango_header = obtener_rango_combinado(ws, fila_header, ultima_columna)

                    if rango_header and rango_header.min_row == rango_header.max_row:
                        col_inicio = rango_header.min_col
                        col_fin = rango_header.max_col
                    else:
                        col_inicio = ultima_columna
                        col_fin = ultima_columna

                    cantidad_columnas = col_fin - col_inicio + 1
                    nueva_col_inicio = col_fin + 1

                    # FECHA BASE
                    header_base = ws.cell(row=fila_header, column=col_inicio)
                    fecha = header_base.value

                    if isinstance(fecha, str):
                        try:
                            fecha = datetime.strptime(fecha, "%d/%m/%Y")
                        except Exception:
                            fecha = None
                    elif isinstance(fecha, (int, float)):
                        from openpyxl.utils.datetime import from_excel
                        fecha = from_excel(fecha)

                    if not fecha:
                        continue

                    nueva_fecha = fecha + relativedelta(months=1)
                    if fecha_archivo is None:
                        fecha_archivo = nueva_fecha

                    # COPIAR HEADER COMPLETO
                    for i in range(cantidad_columnas):
                        col_origen = col_inicio + i
                        col_destino = nueva_col_inicio + i

                        header_origen = ws.cell(row=fila_header, column=col_origen)
                        header_destino = ws.cell(row=fila_header, column=col_destino)

                        header_destino.value = header_origen.value
                        header_destino.font = copy(header_origen.font)
                        header_destino.border = copy(header_origen.border)
                        header_destino.fill = copy(header_origen.fill)
                        header_destino.alignment = copy(header_origen.alignment)

                    ws.cell(row=fila_header, column=nueva_col_inicio).value = nueva_fecha
                    ws.cell(row=fila_header, column=nueva_col_inicio).number_format = "DD/MM/YYYY"

                    if cantidad_columnas > 1:
                        ws.merge_cells(
                            start_row=fila_header,
                            end_row=fila_header,
                            start_column=nueva_col_inicio,
                            end_column=nueva_col_inicio + cantidad_columnas - 1
                        )

                    # FILAS ANTERIORES AL HEADER (ej: fila 1 con porcentajes históricos)
                    offset = nueva_col_inicio - col_inicio

                    # Copiar merges de filas anteriores al header
                    merges_pre_header = [
                        r for r in list(ws.merged_cells.ranges)
                        if r.min_row < fila_header
                        and r.min_col >= col_inicio
                        and r.max_col <= col_fin
                    ]
                    for rango in merges_pre_header:
                        try:
                            # Deshacer merge destino si ya existe (evita conflicto)
                            ws.unmerge_cells(
                                start_row=rango.min_row,
                                end_row=rango.max_row,
                                start_column=rango.min_col + offset,
                                end_column=rango.max_col + offset
                            )
                        except Exception:
                            pass
                        try:
                            ws.merge_cells(
                                start_row=rango.min_row,
                                end_row=rango.max_row,
                                start_column=rango.min_col + offset,
                                end_column=rango.max_col + offset
                            )
                        except Exception:
                            pass

                    for fila in range(1, fila_header):
                        for i in range(cantidad_columnas):
                            col_origen = col_inicio + i
                            col_destino = nueva_col_inicio + i

                            celda_origen = ws.cell(row=fila, column=col_origen)
                            celda_destino = ws.cell(row=fila, column=col_destino)

                            # Saltar celdas secundarias de un merge (solo escribir en la principal)
                            from openpyxl.cell.cell import MergedCell
                            if isinstance(celda_destino, MergedCell):
                                continue

                            valor = celda_origen.value
                            fmt = celda_origen.number_format or ""

                            if isinstance(valor, (int, float)) and "%" in fmt:
                                celda_destino.value = round(factor - 1, 4)
                            else:
                                celda_destino.value = valor

                            if celda_origen.has_style:
                                celda_destino.font = copy(celda_origen.font)
                                celda_destino.border = copy(celda_origen.border)
                                celda_destino.fill = copy(celda_origen.fill)
                                celda_destino.number_format = celda_origen.number_format
                                celda_destino.alignment = copy(celda_origen.alignment)

                    # DATOS
                    for fila in range(fila_header + 1, ws.max_row + 1):
                        for i in range(cantidad_columnas):
                            col_origen = col_inicio + i
                            col_destino = nueva_col_inicio + i

                            celda_origen = ws.cell(row=fila, column=col_origen)
                            celda_destino = ws.cell(row=fila, column=col_destino)

                            from openpyxl.cell.cell import MergedCell
                            if isinstance(celda_destino, MergedCell):
                                continue

                            valor = celda_origen.value

                            if isinstance(valor, (int, float)):
                                fmt = celda_origen.number_format or ""
                                es_porcentaje = "%" in fmt
                                umbral = 0.20 if es_porcentaje else 20

                                if valor != 0 and valor < umbral:
                                    porcentaje_incremento = factor - 1  # ej: 1.04 → 0.04
                                    if es_porcentaje:
                                        celda_destino.value = round(porcentaje_incremento, 4)
                                    else:
                                        celda_destino.value = round(porcentaje_incremento * 100, 2)
                                else:
                                    celda_destino.value = round(valor * factor, 2)
                            elif isinstance(valor, str):
                                match = re.search(r'(NN\s*[xX]\s*)(\d+)', valor)
                                if match:
                                    prefijo = match.group(1)
                                    numero = match.group(2)
                                    celda_destino.value = f"{prefijo}{round(float(numero) * factor, 2)}"
                                else:
                                    celda_destino.value = valor
                            else:
                                celda_destino.value = valor

                            if celda_origen.has_style:
                                celda_destino.font = copy(celda_origen.font)
                                celda_destino.border = copy(celda_origen.border)
                                celda_destino.fill = copy(celda_origen.fill)
                                celda_destino.number_format = celda_origen.number_format
                                celda_destino.alignment = copy(celda_origen.alignment)

                    # COPIAR CELDAS COMBINADAS DEL BLOQUE
                    merges_datos = [
                        r for r in list(ws.merged_cells.ranges)
                        if r.min_row > fila_header
                        and r.min_col >= col_inicio
                        and r.max_col <= col_fin
                    ]
                    for rango in merges_datos:
                        try:
                            ws.unmerge_cells(
                                start_row=rango.min_row,
                                end_row=rango.max_row,
                                start_column=rango.min_col + offset,
                                end_column=rango.max_col + offset
                            )
                        except Exception:
                            pass
                        try:
                            ws.merge_cells(
                                start_row=rango.min_row,
                                end_row=rango.max_row,
                                start_column=rango.min_col + offset,
                                end_column=rango.max_col + offset
                            )
                        except Exception:
                            pass

                    # ANCHO COLUMNAS
                    for i in range(cantidad_columnas):
                        col_orig = get_column_letter(col_inicio + i)
                        col_dest = get_column_letter(nueva_col_inicio + i)
                        ws.column_dimensions[col_dest].width = ws.column_dimensions[col_orig].width

                # GUARDAR
                nombre_base = os.path.splitext(os.path.basename(ruta))[0]
                carpeta = os.path.dirname(ruta)

                nombre_base = re.sub(
                    r'-(0[1-9]|[12][0-9]|3[01])'
                    r'(0[1-9]|1[0-2])'
                    r'20\d{2}(-\d+)?$',
                    '',
                    nombre_base
                )

                extension_salida = ".xlsx"

                if fecha_archivo:
                    sufijo = fecha_archivo.strftime("%d%m%Y")
                    ruta_nueva = os.path.join(carpeta, f"{nombre_base}-{sufijo}{extension_salida}")
                    contador = 2
                    while os.path.exists(ruta_nueva):
                        ruta_nueva = os.path.join(carpeta, f"{nombre_base}-{sufijo}-{contador}{extension_salida}")
                        contador += 1
                else:
                    contador = 2
                    ruta_nueva = os.path.join(carpeta, f"{nombre_base}-{contador}{extension_salida}")
                    while os.path.exists(ruta_nueva):
                        contador += 1
                        ruta_nueva = os.path.join(carpeta, f"{nombre_base}-{contador}{extension_salida}")

                wb.save(ruta_nueva)
                exitos += 1

            except Exception as e:
                print(f"Error en {ruta}: {e}")

        return exitos


if __name__ == "__main__":
    if not verificar_licencia():
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(
            "Acceso denegado",
            "Esta versión del programa no está habilitada.\n"
            "Contactá al desarrollador."
        )
        sys.exit(0)

    app = App()
    app.mainloop()